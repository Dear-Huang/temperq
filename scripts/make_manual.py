from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "使用手册.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_font(run, name="Aptos", size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc, text, style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        set_font(p.add_run(bold_prefix), bold=True)
        set_font(p.add_run(text[len(bold_prefix):]))
    else:
        set_font(p.add_run(text))
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, name="Consolas", size=9)
    return p


def add_picture(doc, name, caption):
    """Screenshot from docs/img (taken from a clean install by the release check)."""
    path = ROOT / "docs" / "img" / name
    if not path.exists():
        return
    doc.add_picture(str(path), width=Cm(16))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(caption), size=9, color="555555")


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12.5)):
        st = styles[name]
        st.font.name = "Aptos Display"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(p.add_run("qwb 本地量化工作台使用手册"), name="Aptos Display", size=22, bold=True)
    p = doc.add_paragraph()
    set_font(p.add_run("给非研发人员的 Windows 与 Linux 操作说明"), size=11, color="555555")
    add_text(doc, "这份手册说明如何在自己的电脑上准备数据、训练 LightGBM 选股模型并阅读报告。第一次使用先跑试运行，确认流程正常后再使用真实数据。数据、训练和结果都留在本机。")

    sections = [
        ("一、Windows 一键使用", [
            "1. 从压缩包中解压整个 qwb 文件夹，不要只复制 bat 文件。",
            "2. 双击 1_安装.bat。它会在当前文件夹创建独立的 Python 3.11 环境并安装固定版本依赖。第一次需要联网，通常几分钟完成。",
            "3. 安装窗口出现“安装完成”后，双击 2_启动.bat。浏览器打开 http://127.0.0.1:8765。",
            "4. 关闭启动窗口即可停止工作台。浏览器可以关闭，训练任务仍会在本机继续。",
            "安装默认用 CPU 训练 LightGBM，所以 NVIDIA、AMD、核显和没有显卡的电脑都能使用。只有使用 Kronos 等需要 PyTorch 的底座时，才选择 GPU 选项；GPU 选项针对 NVIDIA 驱动和 CUDA 版本。",
        ]),
        ("二、Linux 安装和启动", [
            "在终端执行以下命令：",
        ]),
    ]
    for heading, paras in sections:
        doc.add_paragraph(heading, style="Heading 1")
        for para in paras:
            add_text(doc, para)
        if heading.startswith("二"):
            add_code(doc, "cd /path/to/qwb/quant-workbench\nsh install.sh\nsh start.sh --no-browser")
            add_text(doc, "服务器只监听 127.0.0.1。远程服务器上运行时，在自己的电脑另开终端做端口转发：")
            add_code(doc, "ssh -N -L 8765:127.0.0.1:8765 user@server")
            add_text(doc, "然后在本地浏览器打开 http://127.0.0.1:8765。默认 LightGBM 不依赖 GPU；Debian/Ubuntu 若提示缺少运行库，执行 sudo apt install libgomp1。")

    doc.add_paragraph("三、网页四步流程", style="Heading 1")
    add_picture(doc, "home.png", "图 1  首页：从上到下依次是四个步骤")
    for h, text in [
        ("第 1 步 要打造什么模型：用在哪", "先说清楚模型要用在哪里，后面的数据要求和参数都由这一步决定：市场（目前 A 股）、持有期（1 天 / 5 天 / 20 天）、股票池（全 A 股、剔除 ST、只要 ST、沪深 300 / 中证 500 / 中证 1000 成分股、自选名单）、资金量（0–30 万、30–100 万、100–600 万、600 万–1000 万、1000 万–1 亿、1 亿–5 亿、5 亿及以上）、交易方式（只做多）和模型（LightGBM 或 Ridge）。灰色选项是规划中，暂时不能选。资金量决定持股只数和最低成交额：每只仓位不超过该股近 20 日日均成交额的 5%，资金越大，能选的股票越少。指数成分股按每个历史时点当时的成分取。"),
        ("第 2 步 准备数据", "可以选择试运行数据、本机已有 qwb 数据、上传 CSV，或使用 Tushare 兼容接口下载。试运行数据只用于熟悉流程，几秒完成，结果没有投资含义，也不计入评估台账。"),
        ("第 3 步 参数（由第 1 步推出）", "参数由第 1 步推出，页面会列出“因为你选了什么，所以默认怎样训练”。可以改持股只数和佣金；选 LightGBM 时可改树数量、叶子数、学习率、种子数，选 Ridge 时可改正则强度。改动过的值会列在报告首页，便于复核。最终测试年份会根据数据的最后年份自动选择。"),
        ("第 4 步 开始训练", "点击开始训练。网页只生成 YAML 并提交本地任务，训练逻辑与命令行相同。任务页面显示状态、进度、日志和结果文件夹。"),
    ]:
        doc.add_paragraph(h, style="Heading 2")
        add_text(doc, text)
    add_picture(doc, "result.png", "图 2  结果页：先看“一句话结论”，再看扣成本后的表格")

    doc.add_paragraph("CSV 文件要求", style="Heading 2")
    add_text(doc, "每行是一只股票的一天。必填列如下；网页中的下载模板可以直接保存后填写。建议至少三年数据。")
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ("列", "含义", "例子")):
        cell.text = value
        set_cell_shading(cell, "2F5D8A")
        set_cell_border(cell)
        for run in cell.paragraphs[0].runs:
            set_font(run, size=9.5, bold=True, color="FFFFFF")
    for row in [("code", "股票代码", "000001 或 000001.SZ"), ("date", "交易日", "2024-01-02"),
                ("open/high/low/close", "未复权价格", "9.39"), ("amount", "成交额，单位元", "1158366000")]:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            for run in cell.paragraphs[0].runs:
                set_font(run, size=9.5)
    add_text(doc, "可选列是 adj_factor、is_st、up_limit、down_limit。没有涨跌停价时 qwb 按规则包计算；没有 is_st 时不能做 ST 专题。Tushare token 只在下载期间使用，不写入日志。")

    doc.add_paragraph("四、如何读结果", style="Heading 1")
    add_text(doc, "报告首页第一节是扣成本检验。模型和简单基线采用同一规则：每 5 个交易日买入预测前 10%，按标签开仓价成交，涨停买不进、跌停卖不出继续持有，成本包含 STRUCT 半价差、佣金和卖出印花税。它回答“排序能力能否变成收益”，不代表实盘组合，也没有计入大单冲击成本。")
    add_text(doc, "随后查看 IC 与 Newey–West t 值、多种子相关性和前 10% 重叠、前视检查、简单基线与截面打乱，以及台账中的重复测试提示。试运行结果只说明流程走通，没有投资含义。")

    doc.add_paragraph("五、命令行等价操作", style="Heading 1")
    add_text(doc, "网页保存的配置在 runs/ui_configs/。Windows PowerShell：")
    add_code(doc, ".venv\\Scripts\\python.exe -X utf8 -m qwb.cli run runs\\ui_configs\\my_first_model.yaml")
    add_text(doc, "Linux：")
    add_code(doc, ".venv/bin/python -X utf8 -m qwb.cli run runs/ui_configs/my_first_model.yaml")
    add_text(doc, "只检查配置而不训练，在命令末尾加 --check。训练结果默认在 runs/<名称>/，主要文件是 report.md、metrics.json、predictions.parquet 和模型文件。")

    doc.add_paragraph("六、常见问题", style="Heading 1")
    for text in [
        "浏览器打不开：确认启动窗口仍在运行，并访问 127.0.0.1；端口被占用时使用窗口打印出的新端口。",
        "显卡不同：首版 LightGBM 固定走 CPU，不要为了 LightGBM 安装 PyTorch。",
        "CSV 成交额单位：Tushare 的 amount 通常是千元，导入前换算成元。",
        "内存不足：先跑试运行，再使用较小 CSV 或较少股票。",
        "扣成本检验缺失：安装包中应同时有 quant-workbench 和 ashare-tcost 两个目录；源码运行时把两个仓库放在同一级目录。",
    ]:
        add_text(doc, text, bold_prefix=text.split("：")[0] + "：")

    doc.add_paragraph("七、数据与安全边界", style="Heading 1")
    add_text(doc, "qwb 默认只监听本机，网页提交需要本次进程令牌。不要把含有 Tushare token 的配置或本地数据目录上传到公共仓库。qwb 是研究和复现实验工具，不提供荐股、收益承诺或自动实盘交易功能。")

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("qwb 本地量化工作台"), size=9, color="777777")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
